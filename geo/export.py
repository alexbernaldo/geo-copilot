"""Exportación de la propuesta a Markdown y DOCX con anexo de trazabilidad (F11)."""

from pathlib import Path
import re
import docx
from docx.shared import Pt, RGBColor

from geo.schemas import Propuesta
from geo.text import formato_es


def _build_traceability_annex_md(propuesta: Propuesta) -> str:
    """Construye la tabla Markdown del Anexo de Trazabilidad (hallazgo -> evidencias -> ejecucion_ids)."""
    lines = [
        "| Hallazgo | Métrica | Valor | n | Tabla | Ejecucion_IDs de evidencia |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    for h in propuesta.hallazgos:
        if not h.aprobado:
            continue
        for ev in h.evidencias:
            eids = ", ".join(ev.ejecucion_ids[:6])
            lines.append(
                f"| {h.id} ({h.tipo}) | `{ev.metrica}` | {ev.valor} | {ev.n} | {ev.tabla} | {eids} |"
            )
    return formato_es("\n".join(lines))


def export_md(propuesta: Propuesta, path: Path | None = None) -> str:
    """Convierte una Propuesta en un documento Markdown completo (con índice, sello revisor y anexo)."""
    run_id = str(propuesta.metadatos.get("run_id", "demo"))
    fecha = str(propuesta.metadatos.get("fecha", ""))
    rev = propuesta.revision
    sello = (
        f"Verificador GEO: {rev.cifras_verificadas}/{rev.cifras_verificadas + rev.cifras_no_trazables} "
        f"cifras trazables · {len(rev.incidencias)} incidencias normativas · "
        f"Estado: {'APROBADO' if rev.aprobado else 'REVISIÓN REQUERIDA'}"
    )

    parts: list[str] = [
        f"# Propuesta GEO ({propuesta.modo.upper()}) — {propuesta.marca}",
        f"> **Aviso:** Datos sintéticos del kit UFV-EPS × reload. · `run_id={run_id}` · Fecha: `{fecha}`  \n"
        f"> **{sello}**",
        "## Índice de secciones",
    ]

    for sec_title in propuesta.secciones:
        parts.append(f"- {sec_title}")
    parts.append("- Anexo de trazabilidad (Hallazgos → Evidencias → Ejecuciones)")
    parts.append("---")

    for sec_title, sec_body in propuesta.secciones.items():
        parts.append(f"## {sec_title}\n\n{sec_body}")

    parts.append(
        "## Anexo de trazabilidad (Hallazgos → Evidencias → Ejecuciones)\n\n"
        + _build_traceability_annex_md(propuesta)
    )
    parts.append(
        f"\n---\n_Datos sintéticos · run_id: {run_id} · fecha: {fecha}_"
    )

    md_text = "\n\n".join(parts) + "\n"
    if path is not None:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(md_text, encoding="utf-8")
    return md_text


def _add_markdown_table_to_doc(doc: docx.Document, table_lines: list[str]) -> None:  # type: ignore[valid-type]
    """Convierte un bloque de líneas de tabla Markdown en una tabla nativa de python-docx."""
    rows_cells: list[list[str]] = []
    for ln in table_lines:
        s = ln.strip().strip("|")
        cells = [c.strip() for c in s.split("|")]
        if all(re.fullmatch(r"[-:\s]+", c) for c in cells if c):
            continue
        rows_cells.append(cells)

    if not rows_cells:
        return

    n_cols = max(len(r) for r in rows_cells)
    table = doc.add_table(rows=len(rows_cells), cols=n_cols)
    table.style = "Table Grid"
    for r_idx, r_vals in enumerate(rows_cells):
        row_obj = table.rows[r_idx]
        for c_idx in range(n_cols):
            val = r_vals[c_idx] if c_idx < len(r_vals) else ""
            clean_val = val.replace("**", "").replace("`", "")
            cell = row_obj.cells[c_idx]
            cell.text = clean_val
            if r_idx == 0 and cell.paragraphs and cell.paragraphs[0].runs:
                cell.paragraphs[0].runs[0].bold = True


def _render_markdown_block_to_doc(doc: docx.Document, md_block: str) -> None:  # type: ignore[valid-type]
    """Renderiza párrafos, subtítulos, listas y tablas Markdown dentro de un documento python-docx."""
    lines = md_block.splitlines()
    i = 0
    n = len(lines)
    while i < n:
        ln = lines[i].strip()
        if not ln:
            i += 1
            continue

        # Bloque de tabla Markdown
        if ln.startswith("|") and ln.endswith("|"):
            tbl_lines: list[str] = []
            while i < n and lines[i].strip().startswith("|"):
                tbl_lines.append(lines[i].strip())
                i += 1
            _add_markdown_table_to_doc(doc, tbl_lines)
            doc.add_paragraph("")
            continue

        if ln.startswith("#### "):
            doc.add_heading(ln[5:].strip(), level=3)
        elif ln.startswith("### "):
            doc.add_heading(ln[4:].strip(), level=2)
        elif ln.startswith("- ") or ln.startswith("* "):
            clean_bullet = ln[2:].replace("**", "").replace("`", "")
            doc.add_paragraph(clean_bullet, style="List Bullet")
        else:
            clean_p = ln.replace("**", "").replace("`", "")
            doc.add_paragraph(clean_p)
        i += 1


def export_docx(propuesta: Propuesta, path: Path) -> Path:
    """Genera un archivo Word (.docx) con portada, índice, secciones, tablas formateadas, anexo y pie de página."""
    doc = docx.Document()

    run_id = str(propuesta.metadatos.get("run_id", "demo"))
    fecha = str(propuesta.metadatos.get("fecha", ""))
    rev = propuesta.revision

    # Pie de página con «Datos sintéticos · run_id · fecha»
    for section in doc.sections:
        footer = section.footer
        if footer.paragraphs:
            f_p = footer.paragraphs[0]
        else:
            f_p = footer.add_paragraph()
        f_p.text = f"Datos sintéticos · {run_id} · {fecha}"
        if f_p.runs:
            f_p.runs[0].font.size = Pt(8.5)
            f_p.runs[0].font.color.rgb = RGBColor(100, 100, 100)

    # Portada
    title = doc.add_heading(f"Propuesta Estratégica GEO — {propuesta.marca}", level=0)
    if title.runs:
        title.runs[0].font.color.rgb = RGBColor(20, 45, 90)

    sub_p = doc.add_paragraph(
        f"Modo: {propuesta.modo.capitalize()} · Run ID: {run_id} · Fecha: {fecha}\n"
        f"Sello Revisor: {rev.cifras_verificadas}/{rev.cifras_verificadas + rev.cifras_no_trazables} "
        f"cifras trazables · {len(rev.incidencias)} incidencias normativas\n"
        f"Aviso: Estudio elaborado sobre datos sintéticos del kit UFV-EPS × reload."
    )
    if sub_p.runs:
        sub_p.runs[0].italic = True

    # Índice
    doc.add_heading("Índice de contenidos", level=1)
    for sec_title in propuesta.secciones:
        doc.add_paragraph(sec_title, style="List Bullet")
    doc.add_paragraph(
        "Anexo de trazabilidad (Hallazgos → Evidencias → Ejecuciones)",
        style="List Bullet",
    )

    doc.add_page_break()

    # Secciones principales
    for sec_title, sec_body in propuesta.secciones.items():
        doc.add_heading(sec_title, level=1)
        _render_markdown_block_to_doc(doc, sec_body)

    # Anexo de trazabilidad
    doc.add_heading(
        "Anexo de trazabilidad (Hallazgos → Evidencias → Ejecuciones)", level=1
    )
    _render_markdown_block_to_doc(doc, _build_traceability_annex_md(propuesta))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
