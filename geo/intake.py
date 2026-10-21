"""Extracción del brief de marca (.md, .docx o texto) hacia FichaMarca (F7)."""

from pathlib import Path
import re
from typing import Any
import docx

from geo.config import CIUDADES_CONOCIDAS, PROMPTS_DIR
from geo.llm import complete_json
from geo.schemas import FichaMarca


def read_brief_text(path_or_text: str | Path) -> str:
    """Lee un brief desde ruta (.md, .txt, .docx) o devuelve el texto directamente si no es un archivo."""
    if isinstance(path_or_text, Path) or (
        isinstance(path_or_text, str)
        and "\n" not in path_or_text
        and len(path_or_text) < 400
        and Path(path_or_text).exists()
    ):
        p = Path(path_or_text)
        if p.suffix.lower() == ".pdf":
            import pypdf
            reader = pypdf.PdfReader(str(p))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages)
        if p.suffix.lower() == ".docx":
            doc = docx.Document(str(p))
            lines: list[str] = []
            for para in doc.paragraphs:
                txt = para.text.strip()
                if txt:
                    if txt.startswith("■"):
                        lines.append(f"## {txt.lstrip('■').strip()}")
                    else:
                        lines.append(txt)
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if len(cells) >= 2 and cells[0]:
                        lines.append(f"- **{cells[0]}:** {cells[1]}")
            return "\n".join(lines)
        return p.read_text(encoding="utf-8-sig")
    return str(path_or_text)


def compute_budget_range_from_text(text: str) -> tuple[float | None, float | None]:
    """Calcula en código (presupuesto_min_eur, presupuesto_max_eur) a partir de las cifras del brief.

    Soporta:
    - Presupuesto base + rango porcentual: p. ej. '90.000 €; entre un 20 % y un 30 %' -> (18000.0, 27000.0).
    - Rango directo en euros: p. ej. 'entre 18.000 y 27.000 €' o '18.000 - 27.000 €'.
    - Cifra única en euros: p. ej. '25.000 €'.
    """
    if not text:
        return (None, None)

    # Si se pasa el texto completo del brief, aislar primero la línea de Presupuesto
    pres_line_match = re.search(
        r"\*\*Presupuesto[^*]*:\*\*\s*(.+)", text, flags=re.IGNORECASE
    )
    if pres_line_match:
        text = pres_line_match.group(1).strip()

    def _to_float_eur(num_str: str) -> float:
        # En español el punto suele ser separador de miles (90.000) y la coma decimal
        cleaned = num_str.strip().replace(".", "").replace(",", ".")
        return float(cleaned)

    # Caso 1: cifra base en € + dos porcentajes (ej. 90.000 € ... 20 % y 30 %)
    eur_matches = re.findall(r"(\d{1,3}(?:\.\d{3})+|\d+)\s*€", text)
    pct_matches = re.findall(r"(\d+(?:[.,]\d+)?)\s*%", text)

    if len(eur_matches) == 1 and len(pct_matches) >= 2:
        base_eur = _to_float_eur(eur_matches[0])
        p1 = float(pct_matches[0].replace(",", ".")) / 100.0
        p2 = float(pct_matches[1].replace(",", ".")) / 100.0
        low_p, high_p = min(p1, p2), max(p1, p2)
        return (round(base_eur * low_p, 2), round(base_eur * high_p, 2))

    if len(eur_matches) == 1 and len(pct_matches) == 1:
        base_eur = _to_float_eur(eur_matches[0])
        p = float(pct_matches[0].replace(",", ".")) / 100.0
        val = round(base_eur * p, 2)
        return (val, val)

    # Caso 2: dos cantidades en euros ('18.000 € y 27.000 €' o 'entre 18.000 y 27.000 €')
    range_match = re.search(
        r"(\d{1,3}(?:\.\d{3})+|\d+)\s*(?:€)?\s*(?:y|a|-|–)\s*(\d{1,3}(?:\.\d{3})+|\d+)\s*€",
        text,
        flags=re.IGNORECASE,
    )
    if range_match:
        v1 = _to_float_eur(range_match.group(1))
        v2 = _to_float_eur(range_match.group(2))
        return (min(v1, v2), max(v1, v2))

    if len(eur_matches) >= 2:
        vals = sorted(_to_float_eur(m) for m in eur_matches[:2])
        return (vals[0], vals[1])

    if len(eur_matches) == 1:
        val = _to_float_eur(eur_matches[0])
        return (val, val)

    return (None, None)


def _split_bullet_items(raw: str) -> list[str]:
    """Divide listas separadas por '·', ';' o saltos de línea, limpiando espacios."""
    if not raw:
        return []
    parts = re.split(r"\s*·\s*|\s*;\s*|\n+", raw)
    return [p.strip(" -•\t") for p in parts if p.strip(" -•\t")]


def _clean_competitor_name(raw_item: str) -> str:
    """Elimina comentarios entre paréntesis del nombre de un competidor declarado."""
    cleaned = re.sub(r"\(.*?\)", "", raw_item).strip(" .,;-")
    return cleaned


def parse_template(text: str) -> FichaMarca:
    """Extrae una FichaMarca de forma determinista mediante regex sobre la plantilla del kit (sin LLM)."""
    kv: dict[str, str] = {}
    for m in re.finditer(r"\*\*([^*:]+?):\*\*\s*(.+)", text):
        key = m.group(1).strip().lower()
        val = m.group(2).strip()
        kv[key] = val

    def _get(*keys: str) -> str:
        for k in keys:
            for existing_k, v in kv.items():
                if k in existing_k and v:
                    return v
        return ""

    # Marca
    marca = _get("marca")
    if not marca:
        h_match = re.search(r"#\s*BRIEF DE MARCA\s*[—\-:]\s*(.+)", text, flags=re.IGNORECASE)
        if h_match:
            marca = h_match.group(1).strip().title()

    # Web
    web = _get("web") or None

    # Sector y subsector
    sector_raw = _get("sector y subsector", "sector")
    if "·" in sector_raw:
        parts_sec = [p.strip() for p in sector_raw.split("·", 1)]
        sector = parts_sec[0]
        subsector: str | None = parts_sec[1]
    else:
        sector = sector_raw
        subsector = None

    # Servicios y prioridades
    servicios_raw = _get("servicios", "productos / servicios")
    servicios = _split_bullet_items(servicios_raw)

    prio_raw = _get("prioridades de crecimiento")
    if " y " in prio_raw and "·" not in prio_raw:
        servicios_prioritarios = [p.strip() for p in prio_raw.split(" y ") if p.strip()]
    else:
        servicios_prioritarios = _split_bullet_items(prio_raw)

    diferenciales = _split_bullet_items(_get("diferenciales"))

    # Públicos (sección 3)
    publicos: list[str] = []
    sec3_match = re.search(
        r"##\s*3\..*?\n(.*?)(?=\n##\s*4\.|\Z)", text, flags=re.DOTALL | re.IGNORECASE
    )
    if sec3_match:
        for ln in sec3_match.group(1).splitlines():
            s_ln = ln.strip()
            if s_ln.startswith("- ") and not s_ln.startswith("- **"):
                publicos.append(s_ln[2:].strip())
            elif s_ln and not s_ln.startswith(("_", "Cómo deciden", "- **")):
                publicos.append(s_ln)

    # Ciudades, ciudad principal y ciudad de crecimiento
    mercado_principal_raw = _get("mercado principal", "mercados prioritarios")
    ciudad_crec_raw = _get("prioridad de crecimiento", "ciudades / regiones clave")
    sedes_raw = _get("sedes")

    ciudad_principal: str | None = None
    for c in CIUDADES_CONOCIDAS:
        if re.search(rf"\b{re.escape(c)}\b", mercado_principal_raw, flags=re.IGNORECASE):
            ciudad_principal = c
            break

    ciudad_crecimiento: str | None = None
    for c in CIUDADES_CONOCIDAS:
        if re.search(rf"\b{re.escape(c)}\b", ciudad_crec_raw, flags=re.IGNORECASE):
            ciudad_crecimiento = c
            break

    ciudades: list[str] = []
    geo_text = f"{sedes_raw} {mercado_principal_raw} {ciudad_crec_raw}"
    for c in CIUDADES_CONOCIDAS:
        if re.search(rf"\b{re.escape(c)}\b", geo_text, flags=re.IGNORECASE):
            ciudades.append(c)
    if ciudad_principal and ciudad_principal not in ciudades:
        ciudades.insert(0, ciudad_principal)
    if ciudad_crecimiento and ciudad_crecimiento not in ciudades:
        ciudades.append(ciudad_crecimiento)

    # Competidores declarados
    comp_raw = _get("competidores conocidos")
    competidores_declarados = [
        _clean_competitor_name(item)
        for item in _split_bullet_items(comp_raw)
        if _clean_competitor_name(item)
    ]

    # Situación digital
    situacion_digital: dict[str, str] = {}
    for label, field_key in [
        ("web_y_contenidos", "web y contenidos"),
        ("reputacion", "reputación"),
        ("medios_y_terceros", "medios y terceros"),
        ("tecnico", "técnico"),
        ("herramientas", "herramientas"),
        ("fundacion_y_tamano", "fundación y tamaño"),
        ("sedes", "sedes"),
    ]:
        val = _get(field_key)
        if val:
            situacion_digital[label] = val

    # Hipótesis del cliente
    hip_raw = _get("hipótesis")
    quotes = re.findall(r"[«\"“](.+?)[»\"”]", hip_raw)
    hipotesis_cliente = [q.strip() for q in quotes] if quotes else _split_bullet_items(hip_raw)

    # Objetivos y KPIs
    objetivo_principal = _get("objetivo principal")
    kpis_negocio = _split_bullet_items(_get("kpis"))

    horizonte_raw = _get("horizonte")
    horizonte_meses = 12
    m_meses = re.findall(r"(\d+)\s*meses", horizonte_raw)
    if m_meses:
        horizonte_meses = max(int(x) for x in m_meses)

    # Presupuesto calculado por código
    presupuesto_raw = _get("presupuesto")
    presupuesto_min, presupuesto_max = compute_budget_range_from_text(presupuesto_raw)

    equipo = _get("equipo")
    normativa_raw = _get("normativa")
    restricciones_normativas = [
        s.strip() for s in re.split(r"(?<=\.)\s+|\s*·\s*", normativa_raw) if s.strip()
    ]
    tono = _get("tono")
    audiencia = _get("audiencia")
    formato = _get("formato")

    return FichaMarca(
        marca=marca,
        web=web,
        sector=sector,
        subsector=subsector,
        servicios=servicios,
        servicios_prioritarios=servicios_prioritarios,
        diferenciales=diferenciales,
        publicos=publicos,
        ciudades=ciudades,
        ciudad_principal=ciudad_principal,
        ciudad_crecimiento=ciudad_crecimiento,
        competidores_declarados=competidores_declarados,
        situacion_digital=situacion_digital,
        hipotesis_cliente=hipotesis_cliente,
        objetivo_principal=objetivo_principal,
        kpis_negocio=kpis_negocio,
        horizonte_meses=horizonte_meses,
        presupuesto_min_eur=presupuesto_min,
        presupuesto_max_eur=presupuesto_max,
        equipo=equipo,
        restricciones_normativas=restricciones_normativas,
        tono=tono,
        audiencia=audiencia,
        formato=formato,
    )


def extract_llm(
    text: str,
    *,
    provider: str | None = None,
    model: str | None = None,
    use_cache: bool = True,
) -> FichaMarca:
    """Extrae FichaMarca mediante LLM usando prompts/intake_v1.md, recalculando siempre el presupuesto por código."""
    system_prompt = (PROMPTS_DIR / "intake_v1.md").read_text(encoding="utf-8")
    user_prompt = (
        "Extrae la FichaMarca del siguiente brief:\n"
        f"<datos>\n{text}\n</datos>"
    )
    parsed_fallback = parse_template(text)

    ficha_llm = complete_json(
        system=system_prompt,
        user=user_prompt,
        schema=FichaMarca,
        model=model,
        temperature=0.0,
        prompt_version="intake_v1",
        mock_builder=lambda: parsed_fallback,
        provider=provider,
        use_cache=use_cache,
    )

    # El cálculo del presupuesto SIEMPRE lo realiza el código a partir de los números del texto
    p_min, p_max = compute_budget_range_from_text(text)
    if p_min is not None:
        ficha_llm.presupuesto_min_eur = p_min
    if p_max is not None:
        ficha_llm.presupuesto_max_eur = p_max
    return ficha_llm


def merge(
    parsed: FichaMarca,
    llm_ficha: FichaMarca | None,
    raw_text: str = "",
) -> FichaMarca:
    """Fusiona la ficha extraída por plantilla con la del LLM: el LLM rellena lo que el parser no encontró."""
    if llm_ficha is None:
        merged = parsed.model_copy(deep=True)
    else:
        data_p = parsed.model_dump()
        data_l = llm_ficha.model_dump()
        combined: dict[str, Any] = {}
        for k, val_p in data_p.items():
            val_l = data_l.get(k)
            if val_p in (None, "", [], {}):
                combined[k] = val_l
            else:
                combined[k] = val_p
        merged = FichaMarca(**combined)

    if raw_text:
        p_min, p_max = compute_budget_range_from_text(raw_text)
        if p_min is not None:
            merged.presupuesto_min_eur = p_min
        if p_max is not None:
            merged.presupuesto_max_eur = p_max

    return merged


def find_missing_fields(ficha: FichaMarca) -> list[str]:
    """Identifica qué campos clave de FichaMarca han quedado vacíos para avisar en el Checkpoint 0 de la UI."""
    required_checks: list[tuple[str, object]] = [
        ("marca", ficha.marca),
        ("sector", ficha.sector),
        ("servicios_prioritarios", ficha.servicios_prioritarios),
        ("ciudades", ficha.ciudades),
        ("ciudad_crecimiento", ficha.ciudad_crecimiento),
        ("competidores_declarados", ficha.competidores_declarados),
        ("objetivo_principal", ficha.objetivo_principal),
        ("presupuesto_min_eur", ficha.presupuesto_min_eur),
        ("presupuesto_max_eur", ficha.presupuesto_max_eur),
        ("restricciones_normativas", ficha.restricciones_normativas),
        ("audiencia", ficha.audiencia),
    ]
    faltantes: list[str] = []
    for name, val in required_checks:
        if val is None or val == "" or val == []:
            faltantes.append(name)
    return faltantes


def load_brief(
    path_or_text: str | Path,
    *,
    provider: str | None = None,
    model: str | None = None,
    use_cache: bool = True,
) -> tuple[FichaMarca, list[str]]:
    """Carga un brief (.md, .docx o texto), aplica plantilla + LLM, fusiona y devuelve (FichaMarca, faltantes)."""
    raw_text = read_brief_text(path_or_text)
    parsed = parse_template(raw_text)
    try:
        llm_ficha = extract_llm(
            raw_text, provider=provider, model=model, use_cache=use_cache
        )
    except Exception:
        llm_ficha = None

    ficha = merge(parsed, llm_ficha, raw_text=raw_text)
    faltantes = find_missing_fields(ficha)
    return ficha, faltantes


def ficha_resumen(ficha: FichaMarca) -> str:
    """Devuelve un resumen en Markdown de la FichaMarca para mostrar en la interfaz."""
    pres_str = (
        f"{ficha.presupuesto_min_eur:,.0f} € – {ficha.presupuesto_max_eur:,.0f} €"
        if ficha.presupuesto_min_eur is not None and ficha.presupuesto_max_eur is not None
        else "No especificado"
    )
    lines = [
        f"### {ficha.marca}",
        f"- **Sector:** {ficha.sector}" + (f" ({ficha.subsector})" if ficha.subsector else ""),
        f"- **Web:** `{ficha.web or 'N/D'}`",
        f"- **Ciudades:** {', '.join(ficha.ciudades)} "
        f"(Principal: {ficha.ciudad_principal or 'N/D'} · Crecimiento: {ficha.ciudad_crecimiento or 'N/D'})",
        f"- **Servicios prioritarios:** {', '.join(ficha.servicios_prioritarios)}",
        f"- **Diferenciales:** {'; '.join(ficha.diferenciales)}",
        f"- **Competidores declarados:** {', '.join(ficha.competidores_declarados)}",
        f"- **Objetivo principal:** {ficha.objetivo_principal}",
        f"- **Presupuesto GEO estimado:** {pres_str} (Horizonte: {ficha.horizonte_meses} meses)",
        f"- **Equipo disponible:** {ficha.equipo}",
        f"- **Audiencia del entregable:** {ficha.audiencia}",
    ]
    return "\n".join(lines)
