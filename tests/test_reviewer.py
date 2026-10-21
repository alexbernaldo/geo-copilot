"""Tests del verificador de cifras y guardia normativa (F10)."""

import docx

from geo.config import BRIEF_MD, RUNS_DIR
from geo.detection import detect
from geo.export import export_docx, export_md
from geo.findings import generate_findings
from geo.intake import load_brief
from geo.judge_rules import judge_rules
from geo.loader import load_ejecuciones, load_fuentes, load_marcas
from geo.metrics import build_brand_summary_metrics, bundle
from geo.reviewer import compliance, verify_entities, verify_numbers
from geo.strategy import build_plan
from geo.writer import write_client, write_prospect
from geo.reference import REFERENCE_DIR


def _sample_context() -> tuple[dict, object, object]:
    """Construye un contexto real de cifras (metrics + hallazgos + plan + ficha) para los tests de F10."""
    marcas = load_marcas(REFERENCE_DIR)
    df_ej = load_ejecuciones(REFERENCE_DIR)
    catalogo = load_fuentes(REFERENCE_DIR)
    df_det = detect(df_ej, marcas)
    evals = judge_rules(df_det, df_ej, marcas, catalogo)
    tablas = bundle(evals, df_ej, marcas)
    ficha, _ = load_brief(BRIEF_MD, provider="mock")
    hallazgos = generate_findings(tablas, ficha=ficha)
    plan = build_plan(ficha, hallazgos, tablas, provider="mock")
    metrics_summary = build_brand_summary_metrics(tablas, marca_brief=ficha.marca)
    ctx = {
        "metrics": metrics_summary,
        "hallazgos": hallazgos,
        "plan": plan,
        "ficha": ficha,
    }
    return ctx, plan, ficha


def test_verify_numbers_traceable_39_passes_and_invented_57_fails() -> None:
    """Un texto con '39 %' (redondeo de 0.388 en métricas) pasa; uno con '57 %' inventado falla."""
    ctx, plan, ficha = _sample_context()

    # '39 %' (de 0.3875 -> 38.8% / 39%) y '4,2 %' (Barcelona) están en métricas -> pasa
    ok_text = (
        "Clínica Alborada Visión alcanza una aparición espontánea del 39 % (38,8 % exacto, n=240), "
        "con un 4,2 % en Barcelona y un 54,2 % en Madrid."
    )
    inc_ok, ver_ok, no_traz_ok = verify_numbers(ok_text, ctx)
    assert no_traz_ok == 0
    assert ver_ok >= 4
    assert inc_ok == []

    # '57 %' inventado -> falla con incidencia bloqueante cifra_no_trazable
    bad_text = "Clínica Alborada Visión alcanza una cuota espontánea del 57 % en Barcelona."
    inc_bad, _, no_traz_bad = verify_numbers(bad_text, ctx)
    assert no_traz_bad >= 1
    assert any(i.tipo == "cifra_no_trazable" and i.bloqueante for i in inc_bad)


def test_compliance_promise_and_denigratory_comparison() -> None:
    """'garantizamos resultados' produce promesa_resultado bloqueante; 'mejor que Láser Nítida' produce comparativa_denigratoria."""
    ctx, plan, ficha = _sample_context()

    # 1. 'garantizamos resultados' -> promesa_resultado bloqueante
    rev_promise = compliance(
        "Con este plan garantizamos resultados en cirugía refractiva en 12 meses.",
        plan=plan,
        ficha=ficha,
        contexto_cifras=ctx,
        provider="mock", study_dir=REFERENCE_DIR,
    )
    assert rev_promise.aprobado is False
    assert any(
        i.tipo == "promesa_resultado" and i.bloqueante for i in rev_promise.incidencias
    )

    # 2. 'mejor que Láser Nítida' -> comparativa_denigratoria bloqueante
    rev_denig = compliance(
        "El seguimiento de Alborada es mucho mejor que Láser Nítida en el postoperatorio.",
        plan=plan,
        ficha=ficha,
        contexto_cifras=ctx,
        provider="mock", study_dir=REFERENCE_DIR,
    )
    assert rev_denig.aprobado is False
    assert any(
        i.tipo == "comparativa_denigratoria" and i.bloqueante
        for i in rev_denig.incidencias
    )

    # 3. Texto limpio y trazable + plan válido -> aprobado=True
    clean_text = (
        "Clínica Alborada Visión registra un 38,8 % de aparición espontánea (n=240) "
        "y propone un plan de 12 meses con presupuesto dentro del rango de 18.000 € a 27.000 €."
    )
    rev_clean = compliance(
        clean_text,
        plan=plan,
        ficha=ficha,
        contexto_cifras=ctx,
        provider="mock", study_dir=REFERENCE_DIR,
    )
    assert rev_clean.aprobado is True
    assert rev_clean.cifras_no_trazables == 0


def test_writer_and_export_client_and_prospect_f11() -> None:
    """Comprueba F11: write_client, write_prospect, export_md y export_docx verificando con python-docx."""
    marcas = load_marcas(REFERENCE_DIR)
    df_ej = load_ejecuciones(REFERENCE_DIR)
    catalogo = load_fuentes(REFERENCE_DIR)
    df_det = detect(df_ej, marcas)
    evals = judge_rules(df_det, df_ej, marcas, catalogo)
    tablas = bundle(evals, df_ej, marcas)
    ficha, _ = load_brief(BRIEF_MD, provider="mock")
    hallazgos = generate_findings(tablas, ficha=ficha)
    plan = build_plan(ficha, hallazgos, tablas, provider="mock")

    run_id = "test_f11"
    run_dir = RUNS_DIR / run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    # 1. Modo cliente -> propuesta_cliente.docx y .md
    prop_cliente = write_client(
        ficha,
        tablas,
        hallazgos,
        plan,
        run_id=run_id,
        provider="mock",
        study_dir=REFERENCE_DIR,
    )
    assert prop_cliente.modo == "cliente"
    assert len(prop_cliente.secciones) == 10
    assert prop_cliente.revision.aprobado is True
    assert prop_cliente.revision.cifras_no_trazables == 0

    docx_path = run_dir / "propuesta_cliente.docx"
    md_cliente_path = run_dir / "propuesta_cliente.md"
    export_md(prop_cliente, md_cliente_path)
    export_docx(prop_cliente, docx_path)

    assert docx_path.exists()
    doc = docx.Document(str(docx_path))
    assert len(doc.paragraphs) > 20
    assert len(doc.tables) >= 8
    footer_text = doc.sections[0].footer.paragraphs[0].text
    assert "Datos sintéticos" in footer_text and run_id in footer_text

    # 2. Modo prospecto (Vistalia Oftalmología) -> propuesta_prospecto_vistalia.md
    prop_vistalia = write_prospect(
        "Vistalia Oftalmología",
        tablas,
        marcas_catalogo=marcas,
        run_id=run_id,
        provider="mock",
        study_dir=REFERENCE_DIR,
    )
    assert prop_vistalia.modo == "prospecto"
    assert prop_vistalia.revision.aprobado is True
    assert prop_vistalia.revision.cifras_no_trazables == 0

    md_prosp_path = run_dir / "propuesta_prospecto_vistalia.md"
    md_prosp_text = export_md(prop_vistalia, md_prosp_path)
    assert md_prosp_path.exists()
    assert "Vistalia Oftalmología" in md_prosp_text
    assert "Asunto:" in md_prosp_text


def test_verify_entities_blocks_mirasierra_and_uncataloged_entities() -> None:
    """verify_entities bloquea marcas ('Mirasierra', 'OftalmoPlus', 'EuroLaser', 'Clínica Ocularia'), dominios ('oftalmologia-alborada.example'), ciudades ('Valencia') y perfiles ('miedo quirúrgico') que no están en el kit."""
    ctx, plan, ficha = _sample_context()

    # 1. Texto válido con entidades del kit -> 0 incidencias
    valid_text = (
        "Clínica Alborada Visión (alboradavision.example) compite en Madrid y Barcelona con "
        "Grupo Hospitalario Valdecruz (valdecruz.example), Instituto Oftálmico Tramuntana "
        "y Vistalia Oftalmología, según rankings como guiaoftalmologia.example."
    )
    assert verify_entities(valid_text, ficha=ficha, study_dir=REFERENCE_DIR) == []

    # 2. Falla explícitamente con "Mirasierra" (tanto en 'Instituto Oftalmológico Mirasierra' como sola)
    inc_mirasierra = verify_entities(
        "El segundo actor es Instituto Oftalmológico Mirasierra (o Mirasierra).",
        ficha=ficha, study_dir=REFERENCE_DIR,
    )
    assert len(inc_mirasierra) >= 1
    assert all(i.bloqueante for i in inc_mirasierra)
    assert any("mirasierra" in i.texto.lower() for i in inc_mirasierra)

    rev_mirasierra = compliance(
        "Clínica Alborada Visión compite con Mirasierra en Madrid.",
        plan=plan,
        ficha=ficha,
        contexto_cifras=ctx,
        provider="mock", study_dir=REFERENCE_DIR,
    )
    assert rev_mirasierra.aprobado is False

    # 3. Falla con OftalmoPlus, EuroLaser, Clínica Ocularia, oftalmologia-alborada.example, miedo quirúrgico y Valencia
    for bad_snippet in (
        "Se detectaron OftalmoPlus y EuroLaser.",
        "La Clínica Ocularia aparece en las respuestas.",
        "El dominio propio es oftalmologia-alborada.example.",
        "Falta contenido para el perfil de paciente con miedo quirúrgico.",
        "Crecimiento previsto en la clínica de Valencia.",
    ):
        incs = verify_entities(bad_snippet, ficha=ficha, study_dir=REFERENCE_DIR)
        assert len(incs) >= 1, f"Debería bloquear: {bad_snippet}"
        assert all(i.bloqueante for i in incs)

